"""Small, real Office/PDF bytes for extraction and native ingestion tests."""
from io import BytesIO


def document_bytes(kind: str) -> tuple[bytes, list[str]]:
    output = BytesIO()
    if kind == "docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph("Before the delivery table")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Owner"
        table.cell(0, 1).text = "Report"
        table.cell(1, 0).text = "WORDTABLEORION"
        table.cell(1, 1).text = "Delivery requires seven reports"
        nested = table.cell(1, 1).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "WORDNESTEDVEGA"
        doc.add_paragraph("After the delivery table")
        doc.save(output)
        markers = ["WORDTABLEORION", "WORDNESTEDVEGA"]
    elif kind == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        nested = group.shapes.add_group_shape()
        nested.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "PPTGROUPORION"
        table = slide.shapes.add_table(1, 2, 0, Inches(2), Inches(6), Inches(1)).table
        table.cell(0, 0).text = "PPTTABLEVEGA"
        table.cell(0, 1).text = "Delivery requires seven reports"
        slide.notes_slide.notes_text_frame.text = "PPTNOTESLYRA"
        deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(output)
        markers = ["PPTGROUPORION", "PPTTABLEVEGA", "PPTNOTESLYRA"]
    elif kind == "pdf":
        from reportlab.pdfgen.canvas import Canvas
        canvas = Canvas(output)
        for page in range(1, 17):
            canvas.drawString(50, 750, f"Page {page}: PDFTAILORION" if page == 16 else f"Background page {page}")
            canvas.showPage()
        canvas.save()
        markers = ["PDFTAILORION"]
    elif kind == "xlsx":
        from openpyxl import Workbook
        book = Workbook()
        book.active.cell(201, 1, "XLSXTAILORION")
        book.save(output)
        book.close()
        markers = ["XLSXTAILORION"]
    else:
        raise ValueError(kind)
    return output.getvalue(), markers
