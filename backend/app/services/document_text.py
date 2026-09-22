from __future__ import annotations

from pathlib import Path

try:
    import pdfplumber as _pdfplumber

    _HAS_PDF = True
except ImportError:
    _HAS_PDF = False

try:
    from docx import Document as _DocxDocument
    from docx.table import Table as _DocxTable
    from docx.text.paragraph import Paragraph as _DocxParagraph

    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False

try:
    from pptx import Presentation as _Presentation

    _HAS_PPTX = True
except ImportError:
    _HAS_PPTX = False

try:
    import openpyxl as _openpyxl

    _HAS_XLSX = True
except ImportError:
    _HAS_XLSX = False


class DocumentExtractionError(ValueError):
    """A content-free failure that must not be indexed as document text."""


def _word_blocks(element, parent):
    """Keep body order and nested tables; merged cells appear only once."""
    for child in element:
        kind = child.tag.rsplit("}", 1)[-1]
        if kind == "p":
            yield _DocxParagraph(child, parent).text
        elif kind == "tbl":
            seen = set()
            for row in _DocxTable(child, parent).rows:
                cells = []
                for cell in row.cells:
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    cells.append("\n".join(_word_blocks(cell._tc, cell)).strip())
                yield "\t".join(cells)


def _slide_text(shapes):
    for shape in shapes:
        if hasattr(shape, "shapes"):
            yield from _slide_text(shape.shapes)
        if shape.has_text_frame and shape.text.strip():
            yield shape.text.strip()
        if shape.has_table:
            for row in shape.table.rows:
                yield "\t".join(cell.text for cell in row.cells if not cell.is_spanned)


def extract_text_from_file(
    path: Path,
    file_type: str,
    *,
    max_chars: int = 4000,
    empty_placeholder: str = "",
    unsupported_placeholder: str = "",
    error_prefix: str = "",
    require_complete: bool = False,
) -> str:
    """Extract a bounded preview, or complete supported text for ingestion.

    Complete mode removes preview page/row limits and fails instead of indexing
    a truncated body or an error placeholder. It does not provide OCR.
    """
    if not path.is_file():
        if require_complete:
            raise DocumentExtractionError("The source file is not available.")
        return "[File not found]"

    try:
        ft = file_type.lower()
        if ft == "pdf" and _HAS_PDF:
            with _pdfplumber.open(path) as pdf:
                pages = [page.extract_text() or "" for page in (pdf.pages if require_complete else pdf.pages[:15])]
            if require_complete and any(page.strip() for page in pages):
                text = "\n\n".join(f"[Page {index + 1}]\n{page}" for index, page in enumerate(pages))
            else:
                text = "\n".join(pages)
        elif ft == "docx" and _HAS_DOCX:
            doc = _DocxDocument(str(path))
            text = "\n".join(_word_blocks(doc.element.body, doc))
        elif ft == "pptx" and _HAS_PPTX:
            parts = []
            presentation = _Presentation(str(path))
            for index, slide in enumerate(presentation.slides):
                slide_texts = [value for value in _slide_text(slide.shapes) if value.strip()]
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame
                    if notes is not None and notes.text.strip():
                        slide_texts.append("[Notes]\n" + notes.text.strip())
                if slide_texts or require_complete:
                    parts.append(f"[Slide {index + 1}]\n" + "\n".join(slide_texts))
            # Page/slide labels alone are not extractable source evidence.
            text = "\n\n".join(parts) if any(part.split("\n", 1)[1].strip() for part in parts) else ""
        elif ft in ("xlsx", "xls") and _HAS_XLSX:
            workbook = _openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            try:
                for sheet in workbook.worksheets:
                    rows = []
                    for row in sheet.iter_rows(max_row=None if require_complete else 200, values_only=True):
                        cells = [str(cell) if cell is not None else "" for cell in row]
                        if any(cell.strip() for cell in cells):
                            rows.append("\t".join(cells))
                    if rows:
                        parts.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
            finally:
                workbook.close()
            text = "\n\n".join(parts)
        elif ft in ("txt", "md", "csv", "json"):
            text = path.read_text(encoding="utf-8", errors="replace")
        else:
            if require_complete:
                raise DocumentExtractionError("This document type or its text parser is not available.")
            return unsupported_placeholder

        text = text.strip()
        if not text:
            return empty_placeholder
        if len(text) > max_chars:
            if require_complete:
                raise DocumentExtractionError(
                    f"Document text exceeds the {max_chars} character indexing limit; split the source document."
                )
            return text[:max_chars] + "\n…[truncated]"
        return text
    except DocumentExtractionError:
        raise
    except Exception as exc:
        if require_complete:
            raise DocumentExtractionError("The document could not be parsed; check its format and integrity.") from None
        return f"{error_prefix}{exc}" if error_prefix else ""
