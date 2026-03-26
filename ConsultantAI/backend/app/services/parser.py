"""Document text extraction — PDF, DOCX, XLSX."""
from __future__ import annotations

from pathlib import Path
from typing import Union


def extract_text(path: Union[str, Path]) -> str:
    p = Path(path)
    suffix = p.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(p)
    elif suffix in (".docx", ".doc"):
        return _extract_docx(p)
    elif suffix in (".xlsx", ".xls"):
        return _extract_xlsx(p)
    elif suffix == ".txt":
        return p.read_text(errors="ignore")
    else:
        return ""


def _extract_pdf(path: Path) -> str:
    import pdfplumber
    texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                texts.append(text)
    return "\n\n".join(texts)


def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            text = "\t".join(str(c) for c in row if c is not None)
            if text.strip():
                rows.append(text)
    return "\n".join(rows)
